"""Apply reviewer decisions from the review workbook back into ``staging.*`` (LLD-REL-02).

The reviewed xlsx (produced by ``release export``) carries a ``reviewer_decision``
column per sheet: ``approve`` / ``edit`` / ``reject`` / blank. This reads those
decisions back and applies them to the matching ``staging.<table>`` rows,
matched by **sheet + natural id + release candidate**:

* ``approve`` → ``review_status = 'approved'``
* ``edit``    → ``review_status = 'edited'`` **and the edited value columns are
  written back** (family_id, units, numerics, arrays, needs_family — the
  facts-content columns; jsonb ``attributes``/``detail`` are lossy in the export
  display and are left as staged)
* ``reject``  → ``review_status = 'rejected'``
* blank       → left ``pending`` (excluded from promote, kept in staging)

A rejected ``capability_row`` cascades: its ``capability_gas`` children are
rejected too. ``capability_gas`` has no editable value columns (just ``gas``), so
``edit`` on a gas row is decision-only.

Validation runs **before any write**: decisions must be one of the three values,
and any edited ``family_id`` must exist in the frozen ``families.yaml``. On any
error nothing is written. The RC advances to ``imported``. Runs on the caller's
``Connection`` and never commits.
"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import Connection, text

from agentkit.extract.families import load_family_ids
from agentkit.release.candidate import set_status

# Natural-id column(s) per sheet.
_NATURAL_ID: dict[str, tuple[str, ...]] = {
    "document": ("doc_id",),
    "product_family": ("family_id",),
    "product": ("product_id",),
    "capability_row": ("cap_id",),
    "capability_gas": ("cap_id", "gas"),
    "company_fact": ("fact_id",),
    "office": ("office_id",),
}

# Value columns written back on `edit` (facts-content + needs_family; jsonb skipped).
_WRITABLE: dict[str, list[str]] = {
    "document": ["kind", "title", "url", "division", "sha256", "page_count", "needs_family"],
    "product_family": ["division", "category", "name", "summary",
                       "applications", "standards", "needs_family"],
    "product": ["family_id", "model_name", "variant", "description", "needs_family"],
    "capability_row": ["family_id", "comp_type", "lubricated", "cooling",
                       "capacity_min", "capacity_max", "capacity_unit",
                       "discharge_p_min", "discharge_p_max", "pressure_unit",
                       "driver", "standards", "needs_family"],
    "capability_gas": [],
    "company_fact": ["kind", "value", "needs_family"],
    "office": ["name", "city", "region", "address", "phone", "email",
               "serves_divisions", "needs_family"],
}

_ARRAY_COLS = {"driver", "standards", "serves_divisions", "applications"}
_NUM_COLS = {"capacity_min", "capacity_max", "discharge_p_min", "discharge_p_max"}
_INT_COLS = {"page_count"}
_BOOL_COLS = {"lubricated", "needs_family"}

_DECISION_COL = "reviewer_decision"
_VALID_DECISIONS = {"approve", "edit", "reject"}
_STATUS = {"approve": "approved", "edit": "edited", "reject": "rejected"}


def _norm_decision(value: object) -> str:
    if value is None:
        return ""
    s = str(value).strip().lower()
    return s


def _to_bool(value: object) -> bool | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return value
    s = str(value).strip().lower()
    if s in ("true", "1", "yes", "y"):
        return True
    if s in ("false", "0", "no", "n"):
        return False
    return None


def _to_array(value: object) -> list[str]:
    if value is None or value == "":
        return []
    if isinstance(value, (list, tuple)):
        return [str(v) for v in value]
    return [p.strip() for p in str(value).split(",") if p.strip()]


def _coerce(col: str, value: object) -> object:
    if col in _ARRAY_COLS:
        return _to_array(value)
    if col in _BOOL_COLS:
        return _to_bool(value)
    if col in _NUM_COLS:
        return None if value in (None, "") else float(value)
    if col in _INT_COLS:
        return None if value in (None, "") else int(value)
    if value == "":
        return None
    return value


class _Decision:
    __slots__ = ("table", "keys", "decision", "edits", "excel_row")

    def __init__(self, table: str, keys: dict, decision: str, edits: dict, excel_row: int):
        self.table = table
        self.keys = keys
        self.decision = decision
        self.edits = edits
        self.excel_row = excel_row


def _read_workbook(xlsx_path: Path) -> list[_Decision]:
    wb = load_workbook(xlsx_path)
    decisions: list[_Decision] = []
    for table, id_cols in _NATURAL_ID.items():
        if table not in wb.sheetnames:
            continue
        ws = wb[table]
        headers = [c.value for c in ws[1]]
        idx = {h: i for i, h in enumerate(headers) if h is not None}
        if _DECISION_COL not in idx:
            continue
        for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            decision = _norm_decision(row[idx[_DECISION_COL]])
            if decision == "":
                # blank → stays pending; still record for counting.
                keys = {c: row[idx[c]] for c in id_cols}
                if all(v is None for v in keys.values()):
                    continue  # trailing empty row
                decisions.append(_Decision(table, keys, "", {}, excel_row))
                continue
            keys = {c: row[idx[c]] for c in id_cols}
            edits = {}
            if decision == "edit":
                for col in _WRITABLE[table]:
                    if col in idx:
                        edits[col] = _coerce(col, row[idx[col]])
            decisions.append(_Decision(table, keys, decision, edits, excel_row))
    return decisions


def _existing_ids(conn: Connection, rc_id: str) -> dict[str, set[tuple]]:
    out: dict[str, set[tuple]] = {}
    for table, id_cols in _NATURAL_ID.items():
        cols = ", ".join(id_cols)
        rows = conn.execute(
            text(f"SELECT {cols} FROM staging.{table} WHERE release_candidate_id = :rc"),
            {"rc": rc_id},
        ).all()
        out[table] = {tuple(r) for r in rows}
    return out


def _validate(decisions: list[_Decision], existing: dict[str, set[tuple]],
              family_ids: set[str]) -> None:
    errors: list[str] = []
    for d in decisions:
        where = f"{d.table} row {d.excel_row} ({d.keys})"
        if d.decision and d.decision not in _VALID_DECISIONS:
            errors.append(f"{where}: unknown decision {d.decision!r}")
            continue
        if d.decision == "":
            continue
        key_tuple = tuple(d.keys[c] for c in _NATURAL_ID[d.table])
        if key_tuple not in existing[d.table]:
            errors.append(f"{where}: no matching staging row")
        if d.decision == "edit" and d.table in ("product", "capability_row"):
            fid = d.edits.get("family_id")
            if fid and fid not in family_ids:
                errors.append(f"{where}: edited family_id {fid!r} not in families.yaml")
    if errors:
        raise ValueError("import validation failed:\n  " + "\n  ".join(errors))


def _apply(conn: Connection, rc_id: str, d: _Decision, *, reviewer: str,
           reviewed_at: datetime) -> None:
    id_cols = _NATURAL_ID[d.table]
    where = " AND ".join(f"{c} = :k_{c}" for c in id_cols)
    params: dict[str, object] = {"rc": rc_id, "rev": reviewer, "at": reviewed_at}
    params.update({f"k_{c}": d.keys[c] for c in id_cols})

    sets = ["review_status = :status", "reviewer = :rev", "reviewed_at = :at"]
    params["status"] = _STATUS[d.decision]
    if d.decision == "edit":
        for col, value in d.edits.items():
            if col in _ARRAY_COLS:
                sets.append(f"{col} = CAST(:v_{col} AS text[])")
            else:
                sets.append(f"{col} = :v_{col}")
            params[f"v_{col}"] = value
    sql = (f"UPDATE staging.{d.table} SET {', '.join(sets)} "
           f"WHERE release_candidate_id = :rc AND {where}")
    conn.execute(text(sql), params)


def import_review(
    conn: Connection,
    rc_id: str,
    xlsx_path: Path,
    *,
    client: str,
    reviewer: str = "review-import",
    reviewed_at: datetime | None = None,
) -> dict[str, dict[str, int]]:
    """Apply the reviewed workbook to ``staging.*`` and advance the RC to ``imported``.

    Returns per-sheet decision counts ``{table: {approve, edit, reject, blank}}``.
    """
    reviewed_at = reviewed_at or datetime.now(timezone.utc)
    family_ids = load_family_ids(client)

    decisions = _read_workbook(Path(xlsx_path))
    existing = _existing_ids(conn, rc_id)
    _validate(decisions, existing, family_ids)

    counts: dict[str, dict[str, int]] = {
        t: {"approve": 0, "edit": 0, "reject": 0, "blank": 0} for t in _NATURAL_ID
    }
    for d in decisions:
        bucket = d.decision if d.decision else "blank"
        counts[d.table][bucket] += 1
        if d.decision:
            _apply(conn, rc_id, d, reviewer=reviewer, reviewed_at=reviewed_at)

    # Cascade: children of a rejected capability_row are rejected too.
    conn.execute(
        text(
            "UPDATE staging.capability_gas g SET review_status = 'rejected', "
            "reviewer = :rev, reviewed_at = :at "
            "WHERE g.release_candidate_id = :rc AND EXISTS ("
            "  SELECT 1 FROM staging.capability_row c "
            "  WHERE c.release_candidate_id = :rc AND c.cap_id = g.cap_id "
            "    AND c.review_status = 'rejected')"
        ),
        {"rc": rc_id, "rev": reviewer, "at": reviewed_at},
    )

    set_status(conn, rc_id, "imported")
    return counts
