"""LLD-REL-02: import reviewer decisions back into staging.

Round-trips through the real export → edit cells → import path on ``seeded_conn``.
Covers approve/edit/reject/blank, edited value write-back, the rejected-parent gas
cascade, reviewer stamping, RC→imported, and pre-write validation.
"""

from __future__ import annotations

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from agentkit.extract.staging_write import write_rows
from agentkit.release.candidate import create_rc, get_rc
from agentkit.release.export import export_rc
from agentkit.release.import_review import import_review

_DOCS = [{"doc_id": "doc.process", "kind": "pdf", "url": "u", "sha256": "s", "page_count": 1}]
_BASE = dict(
    source_doc_id="doc.process", source_locator="p1 §Process",
    section_id="doc.process::s001", confidence=0.9, review_status="pending",
)


def _content() -> list[tuple[str, dict]]:
    keep = ("capability_row", {
        "cap_id": "cap.keep", "family_id": "fam.process_recip", "comp_type": "reciprocating",
        "capacity_max": 20000, "capacity_unit": "Nm3/hr", "needs_family": False,
        "conflict_group": None, "evidence": {"capacity": "Up to 20000 Nm3/hr"}, **_BASE,
    })
    rej = ("capability_row", {
        "cap_id": "cap.rej", "family_id": "fam.oxygen_recip", "comp_type": "reciprocating",
        "capacity_max": 500, "capacity_unit": "Nm3/hr", "needs_family": False,
        "conflict_group": None, "evidence": {}, **_BASE,
    })

    def gas(cap, g):
        return ("capability_gas", {
            "cap_id": cap, "gas": g, "evidence": {"gas": g}, "section_id": "doc.process::s001",
            "confidence": 0.9, "review_status": "pending", "needs_family": False,
            "conflict_group": None,
        })

    prod = ("product", {
        "product_id": "prd.1", "family_id": "fam.process_recip", "model_name": "X",
        "attributes": {"printed": []}, "needs_family": False, "conflict_group": None,
        "evidence": {}, **_BASE,
    })
    fact = ("company_fact", {
        "fact_id": "cf.1", "kind": "certification", "value": "ISO 9001", "detail": {},
        "needs_family": False, "conflict_group": None, "evidence": {}, **_BASE,
    })
    return [keep, rej, gas("cap.keep", "H2"), gas("cap.keep", "N2"),
            gas("cap.rej", "O2"), prod, fact]


def _col(ws, header: str) -> int:
    return [c.value for c in ws[1]].index(header) + 1


def _set(ws, id_cols: dict, column: str, value) -> None:
    """Set `column` on the row whose natural-id cells match id_cols."""
    idxs = {c: _col(ws, c) for c in id_cols}
    for row in range(2, ws.max_row + 1):
        if all(ws.cell(row=row, column=i).value == id_cols[c] for c, i in idxs.items()):
            ws.cell(row=row, column=_col(ws, column)).value = value
            return
    raise AssertionError(f"row {id_cols} not found in {ws.title}")


def _export_and_edit(seeded_conn, tmp_path, edits) -> tuple[str, object]:
    rc = create_rc(seeded_conn, "jyotech")
    write_rows(seeded_conn, rc_id=rc, document_rows=_DOCS, content_rows=_content())
    out = tmp_path / "review.xlsx"
    export_rc(seeded_conn, rc, out)
    wb = load_workbook(out)
    edits(wb)
    wb.save(out)
    return rc, out


def test_import_applies_decisions_edits_and_cascade(seeded_conn, tmp_path) -> None:
    def edits(wb):
        cap = wb["capability_row"]
        # edit: keep row, change family + a numeric
        _set(cap, {"cap_id": "cap.keep"}, "reviewer_decision", "edit")
        _set(cap, {"cap_id": "cap.keep"}, "family_id", "fam.oxygen_recip")
        _set(cap, {"cap_id": "cap.keep"}, "capacity_max", 30000)
        # reject: rej row (its gas child must cascade)
        _set(cap, {"cap_id": "cap.rej"}, "reviewer_decision", "reject")
        gas = wb["capability_gas"]
        _set(gas, {"cap_id": "cap.keep", "gas": "H2"}, "reviewer_decision", "approve")
        _set(gas, {"cap_id": "cap.keep", "gas": "N2"}, "reviewer_decision", "approve")
        # cap.rej / O2 left BLANK on purpose → cascade must still reject it
        wb["product"]  # approve the product
        _set(wb["product"], {"product_id": "prd.1"}, "reviewer_decision", "approve")
        # company_fact left blank → stays pending

    rc, out = _export_and_edit(seeded_conn, tmp_path, edits)
    counts = import_review(seeded_conn, rc, out, client="jyotech")

    def status(table, **k):
        where = " AND ".join(f"{c} = :{c}" for c in k)
        return seeded_conn.execute(
            text(f"SELECT review_status, reviewer, reviewed_at FROM staging.{table} "
                 f"WHERE release_candidate_id = :rc AND {where}"),
            {"rc": rc, **k},
        ).one()

    # edit applied + value write-back
    row = seeded_conn.execute(
        text("SELECT review_status, family_id, capacity_max, reviewer, reviewed_at "
             "FROM staging.capability_row WHERE release_candidate_id=:rc AND cap_id='cap.keep'"),
        {"rc": rc},
    ).one()
    assert row.review_status == "edited"
    assert row.family_id == "fam.oxygen_recip"
    assert row.capacity_max == 30000
    assert row.reviewer == "review-import"
    assert row.reviewed_at is not None

    assert status("capability_row", cap_id="cap.rej").review_status == "rejected"
    # cascade: blank-decision gas child of the rejected parent is rejected
    assert status("capability_gas", cap_id="cap.rej", gas="O2").review_status == "rejected"
    assert status("capability_gas", cap_id="cap.keep", gas="H2").review_status == "approved"
    assert status("product", product_id="prd.1").review_status == "approved"
    # blank → stays pending
    assert status("company_fact", fact_id="cf.1").review_status == "pending"

    assert get_rc(seeded_conn, rc)["status"] == "imported"
    assert counts["capability_row"] == {"approve": 0, "edit": 1, "reject": 1, "blank": 0}
    assert counts["company_fact"]["blank"] == 1


def test_import_rejects_unknown_decision(seeded_conn, tmp_path) -> None:
    def edits(wb):
        _set(wb["capability_row"], {"cap_id": "cap.keep"}, "reviewer_decision", "maybe")
    rc, out = _export_and_edit(seeded_conn, tmp_path, edits)
    with pytest.raises(ValueError, match="unknown decision"):
        import_review(seeded_conn, rc, out, client="jyotech")
    # nothing written: RC stays exported, row stays pending
    assert get_rc(seeded_conn, rc)["status"] == "exported"


def test_import_rejects_edited_family_not_in_yaml(seeded_conn, tmp_path) -> None:
    def edits(wb):
        cap = wb["capability_row"]
        _set(cap, {"cap_id": "cap.keep"}, "reviewer_decision", "edit")
        _set(cap, {"cap_id": "cap.keep"}, "family_id", "fam.does_not_exist")
    rc, out = _export_and_edit(seeded_conn, tmp_path, edits)
    with pytest.raises(ValueError, match="not in families.yaml"):
        import_review(seeded_conn, rc, out, client="jyotech")
    assert get_rc(seeded_conn, rc)["status"] == "exported"
