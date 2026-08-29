"""LLD-REL-01: review xlsx structure — one sheet per table, evidence beside value,
frozen header, decision dropdown, RC advanced to exported. Uses seeded_conn."""

from __future__ import annotations

from openpyxl import load_workbook

from agentkit.extract.staging_write import STAGING_TABLES, write_rows
from agentkit.release.candidate import create_rc, get_rc
from agentkit.release.export import export_rc

_DOCS = [{"doc_id": "doc.process", "kind": "pdf", "url": "u", "sha256": "s", "page_count": 1}]


def _rows() -> list[tuple[str, dict]]:
    base = dict(
        source_doc_id="doc.process", source_locator="p1 §Process",
        section_id="doc.process::s001", confidence=0.9, review_status="pending",
    )
    cap = ("capability_row", {
        "cap_id": "cap.a", "family_id": "fam.process_recip", "comp_type": "reciprocating",
        "capacity_max": 20000, "capacity_unit": "Nm3/hr",
        "needs_family": False, "conflict_group": "cg.fam.process_recip.capacity",
        "evidence": {"capacity": "Up to 20000 Nm3/hr", "comp_type": "Reciprocating"}, **base,
    })
    fact = ("company_fact", {
        "fact_id": "cf.1", "kind": "certification", "value": "ISO 9001:2015",
        "detail": {}, "needs_family": False, "conflict_group": None,
        "evidence": {"value": "ISO 9001:2015", "kind": "certified"}, **base,
    })
    return [cap, fact]


def test_export_workbook_structure(seeded_conn, tmp_path) -> None:
    rc = create_rc(seeded_conn, "jyotech")
    write_rows(seeded_conn, rc_id=rc, document_rows=_DOCS, content_rows=_rows())
    out = tmp_path / "review.xlsx"

    counts = export_rc(seeded_conn, rc, out)
    assert counts["capability_row"] == 1
    assert counts["company_fact"] == 1

    wb = load_workbook(out)
    # One sheet per staging table.
    assert wb.sheetnames == list(STAGING_TABLES)

    cap = wb["capability_row"]
    headers = [c.value for c in cap[1]]
    assert headers[0] == "cap_id"
    assert "capacity_max" in headers
    assert "capacity_max»evidence" in headers  # evidence beside the value
    assert "review_status" in headers
    assert headers[-1] == "reviewer_decision"

    # The evidence quote sits beside its value.
    row2 = {h: c.value for h, c in zip(headers, cap[2])}
    assert row2["capacity_max"] == 20000
    assert row2["capacity_max»evidence"] == "Up to 20000 Nm3/hr"
    assert row2["conflict_group"] == "cg.fam.process_recip.capacity"

    # Frozen header + a decision dropdown are present.
    assert cap.freeze_panes == "A2"
    assert len(cap.data_validations.dataValidation) == 1

    # The RC is advanced to exported.
    assert get_rc(seeded_conn, rc)["status"] == "exported"


_BASE = dict(
    source_doc_id="doc.process", source_locator="p1 §Process",
    section_id="doc.process::s001", confidence=0.9, review_status="pending",
)


def test_empty_attributes_render_blank_not_repr(seeded_conn, tmp_path) -> None:
    # {"printed": []} must export as a blank cell, never the leaked repr "printed=[]".
    rc = create_rc(seeded_conn, "jyotech")
    prod = ("product", {
        "product_id": "prd.1", "family_id": "fam.process_recip", "model_name": "X",
        "attributes": {"printed": []}, "needs_family": False, "conflict_group": None,
        "evidence": {}, **_BASE,
    })
    out = tmp_path / "review.xlsx"
    write_rows(seeded_conn, rc_id=rc, document_rows=_DOCS, content_rows=[prod])
    export_rc(seeded_conn, rc, out)

    ws = load_workbook(out)["product"]
    headers = [c.value for c in ws[1]]
    row2 = {h: c.value for h, c in zip(headers, ws[2])}
    assert row2["attributes"] in (None, "")          # blank, not "printed=[]"


def test_gas_sheet_shows_parent_provenance(seeded_conn, tmp_path) -> None:
    # capability_gas has no source columns of its own; the sheet fills them display-only
    # from the parent capability_row so the reviewer can see where the row came from.
    rc = create_rc(seeded_conn, "jyotech")
    cap = ("capability_row", {
        "cap_id": "cap.a", "family_id": "fam.process_recip",
        "needs_family": False, "conflict_group": None, "evidence": {}, **_BASE,
    })
    gas = ("capability_gas", {
        "cap_id": "cap.a", "gas": "H2", "evidence": {"gas": "H2"},
        "section_id": "doc.process::s001", "confidence": 0.9, "review_status": "pending",
        "needs_family": False, "conflict_group": None,
    })
    out = tmp_path / "review.xlsx"
    write_rows(seeded_conn, rc_id=rc, document_rows=_DOCS, content_rows=[cap, gas])
    export_rc(seeded_conn, rc, out)

    ws = load_workbook(out)["capability_gas"]
    headers = [c.value for c in ws[1]]
    row2 = {h: c.value for h, c in zip(headers, ws[2])}
    assert row2["source_doc_id"] == "doc.process"
    assert row2["source_locator"] == "p1 §Process"
